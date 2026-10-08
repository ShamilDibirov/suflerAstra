"""Real basic-parser fixtures and memory-bounded multilingual model inference."""
import asyncio
import base64
import os
import resource
import tempfile
import time
from pathlib import Path
import sys
sys.path.insert(0, "/app")
import main
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from fastapi import HTTPException

with tempfile.TemporaryDirectory() as directory:
    path = Path(directory) / 'source.docx'
    doc = Document()
    doc.add_paragraph('Замена SIM: проверка владельца номера.')
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = 'Регламент'
    doc.save(path)
    assert 'Замена SIM' in main.basic_extract(path)
    assert 'Регламент' in main.basic_extract(path)
    path = Path(directory) / 'source.pdf'
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 20 200 Td (Approved SIM process) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)
    assert 'Approved SIM process' in main.basic_extract(path)
    writer = PdfWriter(); writer.add_blank_page(width=300, height=300); writer.write(path)
    try:
        main.parse_file(main.ParseRequest(name='scan.pdf',contentBase64=base64.b64encode(path.read_bytes()).decode()))
    except HTTPException as error:
        assert error.status_code == 422
    else:
        raise AssertionError('An empty/scanned PDF must not produce knowledge')
    path = Path(directory) / 'large.pdf'; path.write_bytes(b'0'*(5*1024*1024+1))
    try: main.basic_extract(path)
    except HTTPException as error: assert error.status_code == 413
    else: raise AssertionError('Oversized file accepted')
print('PASS: DOCX text/tables, PDF text, empty PDF refusal, oversized-file refusal.')

# Isolate inference from external Weaviate: actual retrieval is checked separately.
main.ensure_schema = lambda: None
async def inference():
    async with main.lifespan(main.app):
        assert 'rerank' not in main.models and 'parser' not in main.models
        query = main.embedding_text('Как заменить сим-карту?', query=True)
        passages = [main.embedding_text(text, query=False) for text in ['Замена SIM: проверка владельца номера.', 'Продажа чехла для телефона.']]
        started = time.perf_counter()
        vectors = main.models['embed'].encode([query,*passages],normalize_embeddings=True,batch_size=4)
        assert vectors.shape == (3,384)
        assert float(vectors[0] @ vectors[1]) > float(vectors[0] @ vectors[2])
        print(f'PASS: actual multilingual-e5-small Russian embeddings; inference {time.perf_counter()-started:.2f}s; peak RSS {resource.getrusage(resource.RUSAGE_SELF).ru_maxrss/1024:.0f} MiB (emulated test host, not VPS benchmark).')
asyncio.run(inference())
