"""Authenticated document processing, hybrid retrieval and multilingual reranking."""
import base64
import hmac
import os
import tempfile
import threading
import uuid
from contextlib import asynccontextmanager
from pathlib import Path

import httpx
from fastapi import FastAPI, Header, HTTPException, Depends
from pydantic import BaseModel, Field

TOKEN = os.getenv("INTERNAL_SERVICE_TOKEN", "")
WEAVIATE = os.getenv("WEAVIATE_URL", "http://weaviate:8080").rstrip("/")
CLASS = os.getenv("WEAVIATE_CLASS", "SuflerChunk")
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "BAAI/bge-m3")
RERANK_ENABLED = os.getenv("RERANK_ENABLED", "true").lower() == "true"
PARSER_MODE = os.getenv("PARSER_MODE", "docling")
http = httpx.Client(timeout=90, headers={"Authorization": "Bearer " + os.getenv("WEAVIATE_API_KEY", "")})
models = {}
model_lock = threading.Lock()


def authorize(x_internal_token: str = Header(default="")):
    if not TOKEN or not hmac.compare_digest(TOKEN, x_internal_token):
        raise HTTPException(401)


def request(method, path, **kwargs):
    response = http.request(method, WEAVIATE + path, **kwargs)
    response.raise_for_status()
    return response.json() if response.content else {}


def ensure_schema():
    schema = request("GET", "/v1/schema")
    if any(c["class"] == CLASS for c in schema.get("classes", [])):
        return
    request("POST", "/v1/schema", json={"class": CLASS, "vectorizer": "none",
        "vectorIndexType": "hnsw", "properties": [
            {"name": "orgId", "dataType": ["text"], "tokenization": "field"},
            {"name": "documentId", "dataType": ["text"], "tokenization": "field"},
            {"name": "blockId", "dataType": ["text"], "tokenization": "field"},
            {"name": "version", "dataType": ["int"]},
            {"name": "region", "dataType": ["text"], "tokenization": "field"},
            {"name": "title", "dataType": ["text"]}, {"name": "text", "dataType": ["text"]},
        ]})


@asynccontextmanager
async def lifespan(app):
    if len(TOKEN) < 32:
        raise RuntimeError("INTERNAL_SERVICE_TOKEN must have at least 32 characters")
    ensure_schema()
    import torch
    torch.set_num_threads(int(os.getenv("KNOWLEDGE_TORCH_THREADS", "2")))
    from sentence_transformers import SentenceTransformer
    models["embed"] = SentenceTransformer(EMBEDDING_MODEL, device="cpu")
    if RERANK_ENABLED:
        from sentence_transformers import CrossEncoder
        models["rerank"] = CrossEncoder("BAAI/bge-reranker-v2-m3", max_length=1024, device="cpu")
    if PARSER_MODE == "docling":
        from docling.document_converter import DocumentConverter
        models["parser"] = DocumentConverter()
    elif PARSER_MODE != "basic":
        raise RuntimeError("PARSER_MODE must be docling or basic")
    yield
    http.close()


app = FastAPI(lifespan=lifespan, dependencies=[Depends(authorize)])


@app.get("/health")
def health():
    return {"ok": bool(models), "embedding": EMBEDDING_MODEL, "reranker": "BAAI/bge-reranker-v2-m3" if RERANK_ENABLED else "disabled", "parser": PARSER_MODE}


def embedding_text(text, *, query):
    if EMBEDDING_MODEL.startswith("intfloat/multilingual-e5"):
        return ("query: " if query else "passage: ") + text
    return text


def basic_extract(path):
    if path.stat().st_size > 5 * 1024 * 1024:
        raise HTTPException(413, "Basic parser supports files up to 5 MB")
    if path.suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if len(reader.pages) > 100:
            raise HTTPException(413, "Basic parser supports up to 100 PDF pages")
        parts = []
        for page in reader.pages:
            stream = page.get_contents()
            if stream is not None and len(stream.get_data()) > 2 * 1024 * 1024:
                raise HTTPException(413, "PDF page is too complex for the basic parser")
            parts.append(page.extract_text() or "")
        return "\n\n".join(parts)
    import zipfile
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > 25 * 1024 * 1024:
            raise HTTPException(413, "DOCX expanded content is too large")
    from docx import Document
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
    return "\n".join(parts)


class ParseRequest(BaseModel):
    name: str = Field(max_length=255)
    contentBase64: str = Field(max_length=28_000_000)


@app.post("/parse")
def parse_file(payload: ParseRequest):
    suffix = Path(payload.name).suffix.lower()
    if suffix not in (".pdf", ".docx"):
        raise HTTPException(400, "Only PDF and DOCX")
    try:
        data = base64.b64decode(payload.contentBase64, validate=True)
    except ValueError:
        raise HTTPException(400, "Invalid base64")
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(413)
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / ("source" + suffix)
        path.write_bytes(data)
        with model_lock:
            if PARSER_MODE == "basic":
                text = basic_extract(path)
            else:
                result = models["parser"].convert(path)
                text = result.document.export_to_markdown()
        if not text.strip():
            raise HTTPException(422, "No text extracted. Scanned files require external OCR in basic mode.")
        return {"text": text}


class IndexRequest(BaseModel):
    document: dict


@app.post("/index")
def index_document(payload: IndexRequest):
    doc = payload.document
    if not doc.get("orgId") or not doc.get("id") or not doc.get("blocks"):
        raise HTTPException(400)
    blocks = doc["blocks"]
    if len(blocks) > 2000:
        raise HTTPException(400, "Too many chunks")
    with model_lock:
        vectors = models["embed"].encode([embedding_text(b["text"], query=False) for b in blocks], normalize_embeddings=True, batch_size=int(os.getenv("EMBEDDING_BATCH_SIZE", "16"))).tolist()
    objects = [{"class": CLASS, "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f'{doc["orgId"]}:{doc["id"]}:{doc["version"]}:{b["id"]}')),
        "properties": {"orgId": doc["orgId"], "documentId": doc["id"], "blockId": b["id"], "version": doc["version"],
            "region": doc["region"], "title": doc["title"], "text": b["text"]}, "vector": vector}
        for b, vector in zip(blocks, vectors)]
    for offset in range(0, len(objects), 64):
        batch = request("POST", "/v1/batch/objects", json={"objects": objects[offset:offset + 64]})
        if any(item.get("result", {}).get("errors") for item in batch):
            raise HTTPException(502, "Indexing failed; document has not been published")
    return {"indexed": len(objects)}


class SearchRequest(BaseModel):
    orgId: str = Field(min_length=1, max_length=200)
    query: str = Field(max_length=1000)
    region: str = "Все регионы"
    versions: list[dict] = Field(max_length=5000)


@app.post("/search")
def search(payload: SearchRequest):
    import json
    if not payload.versions:
        return {"results": []}
    # All filters are mandatory and constructed by the server. A client cannot supply GraphQL.
    def eq(field, value, numeric=False):
        return '{path:[' + json.dumps(field) + '],operator:Equal,' + ('valueInt:' if numeric else 'valueText:') + json.dumps(value) + '}'
    version_filters = ['{operator:And,operands:[' + eq("documentId", v["id"]) + ',' + eq("version", v["version"], True) + ']}' for v in payload.versions]
    where = '{operator:And,operands:[' + eq("orgId", payload.orgId) + ',{operator:Or,operands:[' + ','.join(version_filters) + ']},{operator:Or,operands:[' + eq("region", "Все регионы") + ',' + eq("region", payload.region) + ']}]}'
    with model_lock:
        vector = models["embed"].encode(embedding_text(payload.query, query=True), normalize_embeddings=True).tolist()
    query = '{Get{' + CLASS + '(hybrid:{query:' + json.dumps(payload.query) + ',vector:' + json.dumps(vector) + ',alpha:0.5},where:' + where + ',limit:20){documentId blockId version title text _additional{score}}}}'
    result = request("POST", "/v1/graphql", json={"query": query})
    if result.get("errors"):
        raise HTTPException(502, "Hybrid retrieval failed")
    candidates = result.get("data", {}).get("Get", {}).get(CLASS, [])
    if not candidates:
        return {"results": []}
    if RERANK_ENABLED:
        with model_lock:
            scores = models["rerank"].predict([(payload.query, item["text"]) for item in candidates])
        ranked = sorted(zip(candidates, scores), key=lambda pair: float(pair[1]), reverse=True)[:5]
    else:
        # Keep Weaviate's hybrid ordering; no invented semantic confidence score.
        ranked = [(item, float(item.get("_additional", {}).get("score") or 0)) for item in candidates[:5]]
    return {"results": [{"documentId": item["documentId"], "blockId": item["blockId"], "version": item["version"], "score": float(score)} for item, score in ranked]}
